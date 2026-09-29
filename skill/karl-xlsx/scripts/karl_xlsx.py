"""KARL Excel helper: recalculate and error-check formulas, inspect workbooks, render QA sheets.

Commands (`python3 -B /mnt/data/karl-xlsx/scripts/karl_xlsx.py <command> …`):

    recalc FILE [--out OUT]   compute every formula with LibreOffice, store the results in the
                              file itself (formatting and charts untouched), report errors as JSON.
                              Replaces FILE unless --out is given; exit 1 on any formula error.
    inspect FILE              sheets, sizes, formulas, merged cells, charts, macros, and what an
                              openpyxl re-save would lose — run before editing a user's workbook
    qa FILE [--pages 1-4]     render to contact sheets under /mnt/data/_karl_qa/ and report problems

From Python: `brand_formats(workbook)`, `series_style(i)` and `brand_chart(chart)` (xlsxwriter)
and `style_header(ws)` (openpyxl) give the Knauf cell and chart styles; invoke karl-brand first.
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import office_qa  # noqa: E402  (shared QA helpers, identical in every office skill)

BRAND_SCRIPTS = os.environ.get('KARL_BRAND_SCRIPTS', '/mnt/data/karl-brand/scripts')
MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
ERRORS = {
    '#NULL!', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#N/A', '#SPILL!', '#CALC!',
    '#GETTING_DATA', '#FIELD!', '#BLOCKED!', '#UNKNOWN!', '#CONNECT!', '#BUSY!',
}
# Excel 2010+ functions must be stored with an _xlfn. prefix, or Excel shows #NAME?.
# openpyxl writes the formula text as given; recalc adds the prefix where it is missing.
NEEDS_XLFN = (
    'AGGREGATE', 'CEILING.MATH', 'CONCAT', 'DAYS', 'FLOOR.MATH', 'FORECAST.LINEAR', 'IFNA',
    'IFS', 'ISOWEEKNUM', 'MAXIFS', 'MINIFS', 'NETWORKDAYS.INTL', 'NORM.DIST', 'NORM.INV',
    'NORM.S.DIST', 'PERCENTILE.EXC', 'PERCENTILE.INC', 'QUARTILE.EXC', 'QUARTILE.INC',
    'RANK.AVG', 'RANK.EQ', 'STDEV.P', 'STDEV.S', 'SWITCH', 'TEXTJOIN', 'VAR.P', 'VAR.S',
    'WORKDAY.INTL', 'XOR',
)
# Dynamic-array and newer functions: the sandbox's LibreOffice cannot compute them and
# openpyxl cannot store them correctly. Precompute in pandas or use INDEX/MATCH instead.
UNSUPPORTED = (
    'BYCOL', 'BYROW', 'CHOOSECOLS', 'CHOOSEROWS', 'DROP', 'EXPAND', 'FILTER', 'HSTACK', 'LAMBDA',
    'LET', 'MAKEARRAY', 'MAP', 'RANDARRAY', 'REDUCE', 'SCAN', 'SEQUENCE', 'SORT', 'SORTBY',
    'TAKE', 'TEXTAFTER', 'TEXTBEFORE', 'TEXTSPLIT', 'TOCOL', 'TOROW', 'UNIQUE', 'VSTACK',
    'WRAPCOLS', 'WRAPROWS', 'XLOOKUP', 'XMATCH',
)


def _fn_re(names):
    return re.compile(r'(?<![\w.])(' + '|'.join(re.escape(n) for n in names) + r')\s*\(', re.IGNORECASE)


XLFN_RE = _fn_re(NEEDS_XLFN)
UNSUPPORTED_RE = _fn_re(UNSUPPORTED)
STRING_RE = re.compile(r'"(?:[^"]|"")*"')


def _outside_strings(formula, pattern):
    return pattern.findall(STRING_RE.sub('""', formula))


def _xml(data):
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
    return etree.fromstring(data, parser)


def _sheet_parts(package):
    """[(sheet name, part name)] in workbook order."""
    workbook = _xml(package.read('xl/workbook.xml'))
    rels = _xml(package.read('xl/_rels/workbook.xml.rels'))
    targets = {r.get('Id'): r.get('Target') for r in rels.iter(f'{{{PKG_REL}}}Relationship')}
    parts = []
    for sheet in workbook.iter(f'{{{MAIN}}}sheet'):
        target = targets.get(sheet.get(f'{{{REL}}}id'), '')
        part = target.lstrip('/') if target.startswith('/') else os.path.normpath('xl/' + target)
        if part in package.namelist():
            parts.append((sheet.get('name'), part.replace('\\', '/')))
    return parts


def _formula_cells(root):
    for cell in root.iter(f'{{{MAIN}}}c'):
        formula = cell.find(f'{{{MAIN}}}f')
        if formula is not None:
            yield cell, formula


def _prefix(formula):
    def add(match):
        return '_xlfn.' + match.group(1).upper() + '('

    pieces, last = [], 0
    for literal in STRING_RE.finditer(formula):
        pieces.append(XLFN_RE.sub(add, formula[last:literal.start()]))
        pieces.append(literal.group(0))
        last = literal.end()
    pieces.append(XLFN_RE.sub(add, formula[last:]))
    return ''.join(pieces)


def _cell_value(cell, value):
    """Store `value` as the formula's cached result, in schema order (f, v, is, extLst)."""
    from lxml import etree

    for old in cell.findall(f'{{{MAIN}}}v') + cell.findall(f'{{{MAIN}}}is'):
        cell.remove(old)
    cell.attrib.pop('t', None)
    if value is None:
        return
    if isinstance(value, bool):
        kind, text = 'b', '1' if value else '0'
    elif isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        from openpyxl.utils.datetime import to_excel

        kind, text = None, format(to_excel(value), '.17g')
    elif isinstance(value, (int, float)):
        kind, text = None, format(value, '.17g') if isinstance(value, float) else str(value)
    elif isinstance(value, str) and value in ERRORS:
        kind, text = 'e', value
    else:
        kind, text = 'str', str(value)
    if kind:
        cell.set('t', kind)
    formula = cell.find(f'{{{MAIN}}}f')
    v = etree.Element(f'{{{MAIN}}}v')
    v.text = text
    formula.addnext(v)


def recalc(path, out=None):
    """Compute all formulas; write their results into the workbook; return a JSON-able report."""
    from lxml import etree
    from openpyxl import load_workbook

    out = out or path
    job = tempfile.mkdtemp(dir=_work())
    report = {'file': out, 'formulas': 0, 'errors': {}, 'error_count': 0, 'prefixed': [], 'unsupported': [], 'empty_results': 0}
    with zipfile.ZipFile(path) as package:
        sheets = _sheet_parts(package)
        roots, cells = {}, {}
        for name, part in sheets:
            root = _xml(package.read(part))
            roots[part] = root
            for cell, formula in _formula_cells(root):
                text = formula.text or ''
                ref = f"'{name}'!{cell.get('r')}"
                report['formulas'] += 1
                for fn in _outside_strings(text, UNSUPPORTED_RE):
                    report['unsupported'].append(f'{ref}: {fn.upper()}')
                fixed = _prefix(text)
                if fixed != text:
                    formula.text = fixed
                    report['prefixed'].append(ref)
                cells.setdefault(name, []).append((cell, part))
        staged = os.path.join(job, 'staged.xlsx')
        _rewrite(package, staged, {p: etree.tostring(r, xml_declaration=True, encoding='UTF-8', standalone=True) for p, r in roots.items()})
    if report['formulas']:
        computed = load_workbook(office_qa.soffice(staged, 'xlsx'), data_only=True)
        for name, entries in cells.items():
            sheet = computed[name] if name in computed.sheetnames else None
            for cell, _ in entries:
                value = sheet[cell.get('r')].value if sheet is not None else None
                _cell_value(cell, value)
                if value is None:
                    report['empty_results'] += 1
                elif isinstance(value, str) and value in ERRORS:
                    report['errors'].setdefault(value, []).append(f"'{name}'!{cell.get('r')}")
                    report['error_count'] += 1
        with zipfile.ZipFile(staged) as package:
            final = os.path.join(job, 'final.xlsx')
            _rewrite(package, final, {p: etree.tostring(r, xml_declaration=True, encoding='UTF-8', standalone=True) for p, r in roots.items()})
        staged = final
    shutil.move(staged, out)
    shutil.rmtree(job, ignore_errors=True)
    report['errors'] = {k: v[:10] + ([f'… {len(v) - 10} more'] if len(v) > 10 else []) for k, v in report['errors'].items()}
    report['prefixed'] = report['prefixed'][:20]
    report['unsupported'] = report['unsupported'][:20]
    report['ok'] = not report['error_count'] and not report['unsupported']
    if report['unsupported']:
        report['hint'] = 'dynamic-array functions do not compute here: precompute in pandas or use INDEX/MATCH'
    return report


def _rewrite(package, target, replaced):
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in package.infolist():
            data = replaced.get(item.filename)
            dst.writestr(item, data if data is not None else package.read(item.filename))


def _work():
    os.makedirs(office_qa.WORK, exist_ok=True)
    return office_qa.WORK


def inspect(path):
    """What is in a workbook, and what editing it with openpyxl would lose."""
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as package:
        names = package.namelist()
    charts = sum(1 for n in names if re.match(r'xl/charts/chart\d+\.xml$', n))
    result = {
        'file': path,
        'macros': 'xl/vbaProject.bin' in names,
        'charts': charts,
        'pivot_tables': sum(1 for n in names if n.startswith('xl/pivotTables/') and n.endswith('.xml')),
        'external_links': sum(1 for n in names if n.startswith('xl/externalLinks/') and n.endswith('.xml')),
        'sheets': [],
        'warnings': [],
    }
    workbook = load_workbook(path, keep_vba=result['macros'])
    result['defined_names'] = list(workbook.defined_names.keys())[:20]
    for ws in workbook.worksheets:
        formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))
        preview = [[_short(c.value) for c in row[:8]] for row in ws.iter_rows(max_row=4)]
        result['sheets'].append({
            'name': ws.title,
            'state': ws.sheet_state,
            'dimensions': ws.dimensions,
            'formulas': formulas,
            'merged': [str(r) for r in list(ws.merged_cells.ranges)[:10]],
            'freeze_panes': ws.freeze_panes,
            'tables': list(ws.tables.keys()),
            'data_validations': len(ws.data_validations.dataValidation),
            'preview': preview,
        })
    if charts:
        result['warnings'].append(f'{charts} chart(s): openpyxl drops charts when it saves a loaded workbook. '
                                  'Rebuild them after editing, or edit the sheet XML instead.')
    if result['external_links']:
        result['warnings'].append('external links: their cached values may be lost on re-save')
    if result['macros']:
        result['warnings'].append('macros: load with keep_vba=True and save as .xlsm')
    return result


def _short(value):
    if value is None:
        return None
    text = str(value)
    return text if len(text) <= 40 else text[:37] + '…'


def _brand():
    if BRAND_SCRIPTS not in sys.path:
        sys.path.insert(0, BRAND_SCRIPTS)
    try:
        import knauf_brand
    except ImportError as exc:
        raise RuntimeError('Invoke the karl-brand skill first (skill tool) for the Knauf styles.') from exc
    return knauf_brand


# A workbook cannot embed fonts, so the brand rule for spreadsheets is Arial throughout.


def brand_formats(workbook):
    """Knauf xlsxwriter formats: header, text, integer, decimal, percent, date, input, total."""
    kb = _brand()
    font = kb.FALLBACK_FAMILY
    base = {'font_name': font, 'font_size': 10, 'font_color': kb.hex('ink.primary'), 'valign': 'top'}

    def fmt(**extra):
        return workbook.add_format({**base, **extra})

    workbook.formats[0].set_font_name(font)  # the workbook default, used by unformatted cells
    workbook.formats[0].set_font_size(10)
    border = {'border': 1, 'border_color': kb.hex('surface.border')}
    return {
        'header': fmt(bold=True, bg_color=kb.hex('surface.neutral'), text_wrap=True, **border),
        'text': fmt(**border),
        'int': fmt(num_format='#,##0', **border),
        'decimal': fmt(num_format='#,##0.00', **border),
        'percent': fmt(num_format='0.0%', **border),
        'date': fmt(num_format='yyyy-mm-dd', **border),
        'input': fmt(font_color=kb.hex('ink.heading'), bg_color=kb.hex('surface.page'), **border),
        'total': fmt(bold=True, top=2, top_color=kb.hex('ink.heading'), num_format='#,##0'),
        'title': fmt(bold=True, font_size=14, font_color=kb.hex('ink.heading')),
    }


def series_style(index):
    """xlsxwriter options for chart series `index` in the Knauf categorical order."""
    kb = _brand()
    colour = kb.token(f'categorical.{index % len(kb.KNAUF["categorical"])}')
    return {'fill': {'color': colour}, 'line': {'color': colour, 'width': 2}, 'border': {'none': True}}


def brand_chart(chart, legend=True):
    """Knauf fonts, inks and gridlines on an xlsxwriter chart (series colours: series_style)."""
    kb = _brand()
    font = {'name': kb.FALLBACK_FAMILY, 'size': 9, 'color': kb.hex('ink.secondary')}
    grid = {'visible': True, 'line': {'color': kb.token('surface.border'), 'width': 0.75}}
    chart.set_x_axis({'num_font': font, 'name_font': font, 'line': {'color': kb.token('surface.border')}})
    y_axis = {'num_font': font, 'name_font': font, 'major_gridlines': grid, 'line': {'none': True}}
    if type(chart).__name__ in ('ChartColumn', 'ChartBar', 'ChartArea'):
        y_axis['min'] = 0  # bar lengths only compare honestly from zero
    chart.set_y_axis(y_axis)
    chart.set_legend({'font': font, 'position': 'bottom'} if legend else {'none': True})
    chart.set_chartarea({'border': {'none': True}})
    chart.set_title({'none': True})
    return chart


def style_header(ws, row=1, freeze=True):
    """Knauf header row for an openpyxl sheet: fill, bold, border, frozen pane, auto-filter."""
    from openpyxl.styles import Border, Font, PatternFill, Side

    kb = _brand()
    font = kb.FALLBACK_FAMILY
    side = Side(style='thin', color=kb.hex('surface.border'))
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name=font, bold=True, color=kb.hex('ink.primary'))
        cell.fill = PatternFill('solid', fgColor=kb.hex('surface.neutral'))
        cell.border = Border(top=side, bottom=side, left=side, right=side)
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)
    if ws.max_column and ws.max_row > row:
        ws.auto_filter.ref = f'A{row}:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}'


def main(argv=None):
    parser = argparse.ArgumentParser(prog='karl_xlsx.py')
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('recalc'); p.add_argument('file'); p.add_argument('--out')
    sub.add_parser('inspect').add_argument('file')
    p = sub.add_parser('qa'); p.add_argument('file'); p.add_argument('--pages')
    args = parser.parse_args(argv)
    if args.command == 'recalc':
        report = recalc(args.file, args.out)
        print(json.dumps(report, indent=1, ensure_ascii=False))
        return 0 if report['ok'] else 1
    if args.command == 'inspect':
        print(json.dumps(inspect(args.file), indent=1, ensure_ascii=False, default=str))
        return 0
    return office_qa.qa(args.file, pages=office_qa.parse_pages(args.pages))


if __name__ == '__main__':
    sys.exit(main())
