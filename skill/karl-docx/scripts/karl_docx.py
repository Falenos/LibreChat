"""KARL Word helper: a Knauf document builder plus editing and QA commands.

Builder (import it; run Python as `python3 -B`):

    import sys; sys.path.insert(0, '/mnt/data/karl-docx/scripts')
    from karl_docx import KnaufDocument
    doc = KnaufDocument(title='Plant audit 2026', subtitle='Findings and actions')
    doc.heading('Summary', 1)
    doc.paragraph('Three of five lines met the target.')
    doc.table([['Line', 'OEE'], ['L1', '82%']], widths_cm=[8, 4])
    doc.save('/mnt/data/plant-audit-2026.docx')

Commands (`python3 -B /mnt/data/karl-docx/scripts/karl_docx.py <command> …`):

    qa FILE [--pages 1-6]                validate, render contact sheets, report fonts
    text FILE                            plain-text dump
    to-pdf FILE [--out PATH]             convert to PDF (the deliverable, not QA)
    merge-runs FILE [-o OUT]             coalesce identically formatted runs
    replace FILE --find T --replace T [--track --author NAME] [-o OUT]
    accept FILE [-o OUT]                 accept every tracked change
    comment FILE --find T --text T [--author NAME] [-o OUT]

Editing commands write to -o, or replace FILE in place when -o is omitted. Inputs outside
/mnt/data (e.g. read-only uploads) must be given an -o.
"""
import argparse
import copy
import datetime
import os
import sys

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import office_qa  # noqa: E402  (shared QA helpers, identical in every office skill)

BRAND_SCRIPTS = os.environ.get('KARL_BRAND_SCRIPTS', '/mnt/data/karl-brand/scripts')
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_NUMERIC = __import__('re').compile(r'^\s*[-+−(]?[€$£]?\s*\d[\d.,\s]*%?\)?\s*$')
XML_NS = 'http://www.w3.org/XML/1998/namespace'


def _w(tag):
    return f'{{{W_NS}}}{tag}'


def _brand():
    if BRAND_SCRIPTS not in sys.path:
        sys.path.insert(0, BRAND_SCRIPTS)
    try:
        import knauf_brand
    except ImportError as exc:
        raise RuntimeError(
            'The Knauf brand helper is not in the sandbox. Invoke the karl-brand skill with '
            'the skill tool (its files carry the tokens, fonts and wordmark), then retry.'
        ) from exc
    return knauf_brand


# ── builder ──────────────────────────────────────────────────────────────────────────────


# Word rejects a package whose elements are out of schema order (LibreOffice does not, so the
# render check would not notice). These are the successors each inserted element must precede.
_SUCCESSORS = {
    'tblBorders': ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription', 'tblPrChange'),
    'shd': ('noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark', 'headers',
            'cellIns', 'cellDel', 'cellMerge', 'tcPrChange'),
    'updateFields': ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids',
                     'mathPr', 'attachedSchema', 'themeFontLang', 'clrSchemeMapping',
                     'doNotIncludeSubdocsInStats', 'doNotAutoCompressPictures', 'forceUpgrade',
                     'captions', 'readModeInkLockDown', 'smartTagType', 'schemaLibrary',
                     'shapeDefaults', 'doNotEmbedSmartTags', 'decimalSymbol', 'listSeparator'),
}


def _insert_ordered(parent, element, name):
    """Insert `element` before its first schema successor present in `parent`."""
    successors = {_w(tag) for tag in _SUCCESSORS[name]}
    for child in parent:
        if child.tag in successors:
            child.addprevious(element)
            return element
    parent.append(element)
    return element


def _set_font(rpr_owner, family):
    """Set every script slot of a run/style font and drop theme overrides that would win."""
    from docx.oxml.ns import qn

    fonts = rpr_owner.get_or_add_rPr().get_or_add_rFonts()  # python-docx keeps schema order
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn(f'w:{attr}'), None)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{attr}'), family)


def _shade(cell, fill_hex):
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn('w:shd'))
    if shd is None:
        shd = _insert_ordered(tcpr, tcpr.makeelement(qn('w:shd'), {}), 'shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)


def _borders(table, colour_hex, size_eighths=6):
    """Every table border in one brand colour (Table Grid's are black)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn('w:tblBorders'))
    if borders is None:
        borders = _insert_ordered(tblpr, OxmlElement('w:tblBorders'), 'tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), str(size_eighths))
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), colour_hex)


def _field(paragraph, instruction, placeholder=''):
    """A complex field (PAGE, NUMPAGES, TOC…) that Word and LibreOffice evaluate."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def char(kind):
        run = OxmlElement('w:r')
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'), kind)
        run.append(fld)
        return run

    paragraph._p.append(char('begin'))
    instr_run = OxmlElement('w:r')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' {instruction} '
    instr_run.append(instr)
    paragraph._p.append(instr_run)
    paragraph._p.append(char('separate'))
    if placeholder:
        paragraph.add_run(placeholder)
    paragraph._p.append(char('end'))


class KnaufDocument:
    """A Word document in Knauf style: brand fonts and inks, wordmark header, numbered footer."""

    def __init__(self, title, subtitle=None, author='', confidentiality='', landscape=False, cover=True):
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt, RGBColor

        self.kb = _brand()
        self.fonts = self.kb.install_fonts()
        self.doc = docx.Document()
        self.title = title
        self._RGB, self._Pt, self._Cm = RGBColor, Pt, Cm
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)  # A4
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        for side in ('left_margin', 'right_margin'):
            setattr(section, side, Cm(2.2))
        section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
        self._styles()
        self._header_footer(confidentiality)
        props = self.doc.core_properties
        props.title, props.author, props.subject = title, author or '', subtitle or ''
        if cover:
            self.doc.add_paragraph(title, style='Title')
            if subtitle:
                self.doc.add_paragraph(subtitle, style='Subtitle')
            meta = ' · '.join(p for p in (author, datetime.date.today().isoformat()) if p)
            if meta:
                self.paragraph(meta, muted=True)

    def _colour(self, token):
        return self._RGB.from_string(self.kb.hex(token))

    def _styles(self):
        Pt = self._Pt
        styles = self.doc.styles
        normal = styles['Normal']
        _set_font(normal.element, self.fonts['body'])
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = self._colour('ink.primary')
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, space_before in (('Title', 26, 0), ('Heading 1', 18, 18), ('Heading 2', 14, 12), ('Heading 3', 12, 10)):
            style = styles[name]
            _set_font(style.element, self.fonts['heading'])
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = self._colour('ink.heading')
            style.paragraph_format.space_before = Pt(space_before)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        # The template's Title carries a theme-coloured bottom rule — not a Knauf token.
        title_ppr = styles['Title'].element.pPr
        if title_ppr is not None:
            for border in title_ppr.findall(_w('pBdr')):
                title_ppr.remove(border)
        subtitle = styles['Subtitle']
        _set_font(subtitle.element, self.fonts['body'])
        subtitle.font.size = Pt(13)
        subtitle.font.italic = False
        subtitle.font.color.rgb = self._colour('ink.secondary')
        footer_style = styles['Footer']
        _set_font(footer_style.element, self.fonts['body'])
        footer_style.font.size = Pt(8.5)
        footer_style.font.color.rgb = self._colour('ink.muted')
        caption = styles['Caption']
        _set_font(caption.element, self.fonts['body'])
        caption.font.size = Pt(9)
        caption.font.italic = False
        caption.font.color.rgb = self._colour('ink.muted')

    def _header_footer(self, confidentiality):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        section = self.doc.sections[0]
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.add_run().add_picture(self.kb.wordmark_png('logo', 240), height=self._Cm(0.9))
        footer = section.footer.paragraphs[0]
        footer.style = self.doc.styles['Footer']  # size and ink come from the style, fields too
        footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
        label = ' · '.join(p for p in (self.title, confidentiality) if p)
        footer.add_run(f'{label}    Page ' if label else 'Page ')
        _field(footer, 'PAGE', '1')
        footer.add_run(' of ')
        _field(footer, 'NUMPAGES', '1')

    def heading(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def paragraph(self, text='', bold_lead=None, muted=False, style=None):
        paragraph = self.doc.add_paragraph(style=style)
        if bold_lead:
            paragraph.add_run(bold_lead).bold = True
            paragraph.add_run(' ')
        run = paragraph.add_run(text)
        if muted:
            run.font.color.rgb = self._colour('ink.muted')
        return paragraph

    def bullets(self, items, numbered=False):
        style = 'List Number' if numbered else 'List Bullet'
        return [self.doc.add_paragraph(item, style=style) for item in items]

    def table(self, rows, header=True, widths_cm=None, align_numbers=True, zebra=True):
        """rows: list of lists of strings. Widths default to an even split of the text width."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        section = self.doc.sections[0]
        usable_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
        ncols = max(len(r) for r in rows)
        widths = widths_cm or [usable_cm / ncols] * ncols
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        table.style = self.doc.styles['Table Grid']
        _borders(table, self.kb.hex('surface.border'))
        body = rows[1:] if header else rows
        numeric_cols = {  # a column whose filled body cells are all numbers aligns right, header too
            c for c in range(ncols)
            if [str(r[c]) for r in body if c < len(r) and str(r[c]).strip()]
            and all(_NUMERIC.match(str(r[c])) for r in body if c < len(r) and str(r[c]).strip())
        }
        for r, row in enumerate(rows):
            for c in range(ncols):
                cell = table.cell(r, c)
                cell.width = self._Cm(widths[c])
                value = str(row[c]) if c < len(row) else ''
                cell.text = value
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = self._Pt(0)
                if align_numbers and (c in numeric_cols or (r > 0 and _NUMERIC.match(value))):
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if header and r == 0:
                    _shade(cell, self.kb.hex('surface.neutral'))
                    for run in para.runs:
                        run.bold = True
                elif zebra and r % 2 == 0:
                    _shade(cell, self.kb.hex('surface.page'))
        self.doc.add_paragraph()
        return table

    def figure(self, image_path, width_cm=16, caption=None):
        self.doc.add_picture(image_path, width=self._Cm(width_cm))
        if caption:
            self.doc.add_paragraph(caption, style='Caption')

    def callout(self, text, title=None):
        """A shaded one-cell box for a key message or a note."""
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        _shade(cell, self.kb.hex('surface.page'))
        cell.text = ''
        if title:
            run = cell.paragraphs[0].add_run(title)
            run.bold = True
            run.font.color.rgb = self._colour('ink.heading')
            cell.add_paragraph(text)
        else:
            cell.paragraphs[0].add_run(text)
        self.doc.add_paragraph()
        return table

    def page_break(self):
        from docx.enum.text import WD_BREAK

        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self, levels='1-3'):
        """A table of contents field. Word fills it when the file opens (it asks to update)."""
        self.heading('Contents', 1)
        _field(self.doc.add_paragraph(), f'TOC \\o "{levels}" \\h \\z \\u', 'Right-click to update the table of contents.')
        _enable_update_fields(self.doc)

    def save(self, path):
        _declare_fonts(self.doc, {self.fonts['body'], self.fonts['heading']} - {self.kb.FALLBACK_FAMILY})
        self.doc.save(path)
        return path


def _declare_fonts(document, families):
    """Font-table entries for the brand families with Arial as the substitute.

    python-docx cannot embed fonts, so a reader without Knauf Screen installed gets a
    substitute. `w:altName` and the swiss family make Word pick Arial, not its default serif.
    """
    from lxml import etree
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    part = next((r.target_part for r in document.part.rels.values() if r.reltype == RT.FONT_TABLE), None)
    if part is None or not families:
        return
    root = etree.fromstring(part.blob)
    present = {f.get(_w('name')) for f in root.findall(_w('font'))}
    for family in sorted(families - present):
        font = etree.SubElement(root, _w('font'))
        font.set(_w('name'), family)
        for tag, value in (('altName', 'Arial'), ('charset', '00'), ('family', 'swiss'), ('pitch', 'variable')):
            etree.SubElement(font, _w(tag)).set(_w('val'), value)
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def _enable_update_fields(document):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = document.settings.element
    if settings.find(qn('w:updateFields')) is None:
        update = OxmlElement('w:updateFields')
        update.set(qn('w:val'), 'true')
        _insert_ordered(settings, update, 'updateFields')


# ── editing ──────────────────────────────────────────────────────────────────────────────


def _paragraphs(document):
    """Every paragraph: body, tables (nested), headers and footers."""
    def walk(container):
        for paragraph in container.paragraphs:
            yield paragraph
        for table in getattr(container, 'tables', []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            yield from walk(part)


def _is_plain_text_run(run_el):
    """Only rPr and w:t children: safe to merge or split without touching fields/drawings."""
    return all(child.tag in (_w('rPr'), _w('t')) for child in run_el)


def _rpr_key(run_el):
    from lxml import etree

    rpr = run_el.find(_w('rPr'))
    return b'' if rpr is None else etree.tostring(rpr, method='c14n')


def _run_text(run_el):
    return ''.join(t.text or '' for t in run_el.findall(_w('t')))


def _set_run_text(run_el, text):
    for t in run_el.findall(_w('t')):
        run_el.remove(t)
    t = run_el.makeelement(_w('t'), {})
    t.text = text
    t.set(f'{{{XML_NS}}}space', 'preserve')
    run_el.append(t)


def merge_runs(document):
    """Coalesce adjacent plain-text runs with identical formatting. Returns the count merged."""
    merged = 0
    for paragraph in _paragraphs(document):
        runs = [child for child in paragraph._p if child.tag == _w('r')]
        previous = None
        for run in runs:
            if previous is not None and run.getprevious() is previous and _is_plain_text_run(run) \
                    and _is_plain_text_run(previous) and _rpr_key(run) == _rpr_key(previous):
                _set_run_text(previous, _run_text(previous) + _run_text(run))
                run.getparent().remove(run)
                merged += 1
                continue
            previous = run
    return merged


def _split(run_el, offset):
    """Split a plain-text run at `offset`; returns the right-hand run (inserted after)."""
    text = _run_text(run_el)
    right = copy.deepcopy(run_el)
    _set_run_text(run_el, text[:offset])
    _set_run_text(right, text[offset:])
    run_el.addnext(right)
    return right


def _isolate(paragraph, start, end):
    """Split runs so [start, end) of the paragraph text is covered by whole runs; return them."""
    position, covered = 0, []
    for run in [child for child in paragraph._p if child.tag == _w('r')]:
        if not _is_plain_text_run(run):
            position += len(_run_text(run))
            continue
        length = len(_run_text(run))
        run_start, run_end = position, position + length
        position = run_end
        if run_end <= start or run_start >= end:
            continue
        if run_start < start:
            run = _split(run, start - run_start)
            run_start = start
        if run_end > end:
            _split(run, end - run_start)
        covered.append(run)
    return covered


_revision_ids = iter(range(9000, 10**9))


def _next_revision(author):
    """Attributes for a w:ins / w:del: a fresh id, the author and a UTC timestamp."""
    return {
        _w('id'): str(next(_revision_ids)),
        _w('author'): author,
        _w('date'): datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
    }


def _tracked_span(find, replacement):
    """The part of `find` that actually changes, as (prefix_len, suffix_len, new_middle)."""
    prefix = len(os.path.commonprefix([find, replacement]))
    suffix = len(os.path.commonprefix([find[prefix:][::-1], replacement[prefix:][::-1]]))
    return prefix, suffix, replacement[prefix:len(replacement) - suffix]


def replace_text(document, find, replacement, track=False, author='KARL', count=None):
    """Replace `find` wherever it occurs within a paragraph, across run boundaries.

    With `track`, only the part that differs is redlined ('80 percent' → '85 percent' marks
    '0' deleted and '5' inserted), the way a reviewer would edit.
    """
    if not find:
        raise ValueError('find must not be empty')
    merge_runs(document)
    done = 0
    for paragraph in _paragraphs(document):
        search_from = 0  # in the coordinates of the paragraph's direct (unrevised) runs
        while count is None or done < count:
            text = ''.join(_run_text(r) for r in paragraph._p if r.tag == _w('r'))
            at = text.find(find, search_from)
            if at < 0:
                break
            if not track:
                runs = _isolate(paragraph, at, at + len(find))
                if not runs or ''.join(_run_text(r) for r in runs) != find:
                    break  # the match spans a field or drawing; leave it alone
                _set_run_text(runs[0], replacement)
                for run in runs[1:]:
                    run.getparent().remove(run)
                search_from = at + len(replacement)
                done += 1
                continue

            prefix, suffix, new_middle = _tracked_span(find, replacement)
            start, end = at + prefix, at + len(find) - suffix
            if start == end and not new_middle:
                search_from = at + len(find)  # identical text: nothing to track
                continue
            if start == end:
                # Pure insertion, anchored after the unchanged prefix. Direct text is unchanged.
                anchor = _isolate(paragraph, start - 1, start) if start > 0 else []
                if not anchor:
                    break
                new_run = copy.deepcopy(anchor[0])
                _set_run_text(new_run, new_middle)
                inserted = anchor[0].makeelement(_w('ins'), _next_revision(author))
                inserted.append(new_run)
                anchor[0].addnext(inserted)
                search_from = at + len(find)
                done += 1
                continue
            runs = _isolate(paragraph, start, end)
            if not runs or ''.join(_run_text(r) for r in runs) != text[start:end]:
                break  # the change spans a field or drawing; leave it alone
            first = runs[0]
            deleted = first.makeelement(_w('del'), _next_revision(author))
            first.addprevious(deleted)
            for run in runs:
                for t in run.findall(_w('t')):
                    t.tag = _w('delText')
                deleted.append(run)
            if new_middle:
                new_run = copy.deepcopy(runs[0])
                for node in new_run.findall(_w('delText')):
                    node.tag = _w('t')
                _set_run_text(new_run, new_middle)
                inserted = deleted.makeelement(_w('ins'), _next_revision(author))
                inserted.append(new_run)
                deleted.addnext(inserted)
            # The deleted run left the direct text; the match now ends that much earlier.
            search_from = at + len(find) - (end - start)
            done += 1
    return done


def accept_all(document):
    """Accept every tracked change: drop deletions, keep insertions, clear format changes."""
    body = document.element
    parts = [body] + [s.header._element for s in document.sections] + [s.footer._element for s in document.sections]
    accepted = 0
    for root in parts:
        for tag in ('del', 'moveFrom'):
            for node in list(root.iter(_w(tag))):
                parent = node.getparent()
                if parent is not None and parent.tag == _w('rPr'):
                    # a deleted paragraph mark: merge this paragraph into the next one
                    paragraph = parent.getparent().getparent()
                    following = paragraph.getnext()
                    if following is not None and following.tag == _w('p'):
                        position = 1 if following.find(_w('pPr')) is not None else 0
                        for child in [c for c in paragraph if c.tag != _w('pPr')]:
                            following.insert(position, child)
                            position += 1
                        paragraph.getparent().remove(paragraph)
                    else:
                        parent.remove(node)
                elif parent is not None:
                    parent.remove(node)
                accepted += 1
        for tag in ('ins', 'moveTo'):
            for node in list(root.iter(_w(tag))):
                parent = node.getparent()
                if parent is None:
                    continue
                if parent.tag == _w('rPr'):
                    parent.remove(node)
                else:
                    index = parent.index(node)
                    for offset, child in enumerate(list(node)):
                        parent.insert(index + offset, child)
                    parent.remove(node)
                accepted += 1
        for tag in ('rPrChange', 'pPrChange', 'sectPrChange', 'tblPrChange', 'trPrChange', 'tcPrChange'):
            for node in list(root.iter(_w(tag))):
                node.getparent().remove(node)
                accepted += 1
    return accepted


def add_comment(document, find, text, author='KARL'):
    """Anchor a comment to the first occurrence of `find`."""
    merge_runs(document)
    for paragraph in _paragraphs(document):
        plain = ''.join(_run_text(r) for r in paragraph._p if r.tag == _w('r'))
        at = plain.find(find)
        if at < 0:
            continue
        run_els = _isolate(paragraph, at, at + len(find))
        runs = [r for r in paragraph.runs if r._r in run_els]
        if runs:
            document.add_comment(runs, text=text, author=author, initials=author[:2].upper())
            return True
    return False


# ── CLI ──────────────────────────────────────────────────────────────────────────────────


def _open(path):
    import docx

    return docx.Document(path)


def _save(document, source, out):
    target = out or source
    if not out and not os.path.abspath(source).startswith('/mnt/data/'):
        raise SystemExit('Give -o: the input is outside /mnt/data and may be read-only.')
    document.save(target)
    print(f'wrote {target}')
    return target


def main(argv=None):
    parser = argparse.ArgumentParser(prog='karl_docx.py')
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('qa'); p.add_argument('file'); p.add_argument('--pages')
    p = sub.add_parser('text'); p.add_argument('file')
    p = sub.add_parser('to-pdf'); p.add_argument('file'); p.add_argument('--out')
    for name in ('merge-runs', 'accept'):
        p = sub.add_parser(name); p.add_argument('file'); p.add_argument('-o', '--out')
    p = sub.add_parser('replace'); p.add_argument('file'); p.add_argument('--find', required=True)
    p.add_argument('--replace', required=True); p.add_argument('--track', action='store_true')
    p.add_argument('--author', default='KARL'); p.add_argument('--count', type=int); p.add_argument('-o', '--out')
    p = sub.add_parser('comment'); p.add_argument('file'); p.add_argument('--find', required=True)
    p.add_argument('--text', required=True); p.add_argument('--author', default='KARL'); p.add_argument('-o', '--out')
    args = parser.parse_args(argv)

    if args.command == 'qa':
        return office_qa.qa(args.file, pages=office_qa.parse_pages(args.pages))
    if args.command == 'text':
        print(office_qa.text_dump(args.file))
        return 0
    if args.command == 'to-pdf':
        office_qa.ensure_fonts()
        pdf = office_qa.soffice(args.file, 'pdf')
        out = args.out or os.path.splitext(args.file)[0] + '.pdf'
        with open(pdf, 'rb') as src, open(out, 'wb') as dst:
            dst.write(src.read())
        print(f'wrote {out}')
        return 0
    document = _open(args.file)
    if args.command == 'merge-runs':
        print(f'merged {merge_runs(document)} runs')
    elif args.command == 'accept':
        print(f'accepted {accept_all(document)} changes')
    elif args.command == 'replace':
        n = replace_text(document, args.find, args.replace, args.track, args.author, args.count)
        print(f'replaced {n} occurrence(s){" as tracked changes" if args.track else ""}')
        if n == 0:
            return 1
    elif args.command == 'comment':
        if not add_comment(document, args.find, args.text, args.author):
            print(f'text not found: {args.find!r}')
            return 1
        print('comment added')
    _save(document, args.file, args.out)
    return 0


if __name__ == '__main__':
    sys.exit(main())
