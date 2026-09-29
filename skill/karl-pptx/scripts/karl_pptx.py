"""KARL PowerPoint helper: a Knauf deck builder plus editing, lint and QA commands.

Builder (import it; run Python as `python3 -B`):

    import sys; sys.path.insert(0, '/mnt/data/karl-pptx/scripts')
    from karl_pptx import KnaufDeck
    deck = KnaufDeck(title='Q3 operations review')
    deck.title_slide('Q3 operations review', 'Plant network · October 2026')
    deck.bullets_slide('What changed', ['OEE up 4 points', 'Scrap down 0.6 points'], notes='…')
    deck.chart_slide('OEE by line', ['L1', 'L2', 'L3'], {'Q2': [80, 76, 82], 'Q3': [84, 79, 85]}, kind='column')
    deck.save('/mnt/data/q3-operations-review.pptx')

Commands (`python3 -B /mnt/data/karl-pptx/scripts/karl_pptx.py <command> …`):

    qa FILE [--pages 1-6]              validate, lint layout, render contact sheets
    lint FILE                          layout lint only (off-slide shapes, overflow, placeholders)
    text FILE                          text of every slide, with slide numbers
    to-pdf FILE [--out PATH]           convert to PDF
    duplicate FILE --slide N [-o OUT]  copy slide N to the end (then reorder)
    delete FILE --slides 2,5 [-o OUT]
    reorder FILE --order 3,1,2,… [-o OUT]   the full new order, every slide once
"""
import argparse
import copy
import json
import os
import subprocess
import sys

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import office_qa  # noqa: E402  (shared QA helpers, identical in every office skill)

BRAND_SCRIPTS = os.environ.get('KARL_BRAND_SCRIPTS', '/mnt/data/karl-brand/scripts')
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5  # 16:9
MARGIN_IN = 0.6
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _brand():
    if BRAND_SCRIPTS not in sys.path:
        sys.path.insert(0, BRAND_SCRIPTS)
    try:
        import knauf_brand
    except ImportError as exc:
        raise RuntimeError(
            'The Knauf brand helper is not in the sandbox. Invoke the karl-brand skill with '
            'the skill tool (its files carry the tokens, fonts and wordmark), then retry.'
        ) from exc
    return knauf_brand


# ── builder ──────────────────────────────────────────────────────────────────────────────


class KnaufDeck:
    """A 16:9 deck in the Knauf slide grammar: blue title/section fields, white content slides."""

    def __init__(self, title, footer_label=None):
        from pptx import Presentation
        from pptx.util import Inches

        self.kb = _brand()
        self.fonts = self.kb.install_fonts()
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
        self.blank = self.prs.slide_layouts[6]
        self.title = title
        self.footer_label = footer_label if footer_label is not None else title
        self._logo = {fill: self.kb.wordmark_png(fill, 300) for fill in ('logo', 'white')}

    # colours and text ------------------------------------------------------------------
    def _rgb(self, token):
        from pptx.dml.color import RGBColor

        return RGBColor.from_string(self.kb.hex(token))

    def _fill(self, slide, token):
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self._rgb(token)

    def _text(self, slide, text, x, y, w, h, size, colour, bold=False, heading=False, align=None,
              anchor='top'):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}[anchor]
        lines = text if isinstance(text, list) else [text]
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            if align:
                paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
            run = paragraph.add_run()
            run.text = line
            font = run.font
            font.size, font.bold = Pt(size), bold
            font.name = self.fonts['heading' if heading else 'body']
            font.color.rgb = self._rgb(colour)
        return box

    def _logo_at(self, slide, fill, x, y, height):
        from pptx.util import Inches

        slide.shapes.add_picture(self._logo[fill], Inches(x), Inches(y), height=Inches(height))

    def _chrome(self, slide, heading):
        """Content-slide frame: heading, footer label, slide number, small wordmark."""
        number = len(self.prs.slides)
        self._text(slide, heading, MARGIN_IN, 0.45, SLIDE_W_IN - 2 * MARGIN_IN - 1.4, 0.9, 28,
                   'ink.heading', bold=True, heading=True, anchor='middle')
        self._text(slide, f'{self.footer_label}', MARGIN_IN, SLIDE_H_IN - 0.45, 8, 0.3, 10, 'ink.muted')
        box = self._text(slide, str(number), SLIDE_W_IN - MARGIN_IN - 1.6, SLIDE_H_IN - 0.45, 0.4, 0.3, 10,
                         'ink.muted', align='right')
        _slide_number_field(box.text_frame.paragraphs[0])  # stays right after reorder/delete
        self._logo_at(slide, 'logo', SLIDE_W_IN - MARGIN_IN - 1.0, SLIDE_H_IN - 0.55, 0.42)

    def _new(self, notes=None):
        slide = self.prs.slides.add_slide(self.blank)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    # slide types -----------------------------------------------------------------------
    def title_slide(self, title, subtitle=None, notes=None):
        slide = self._new(notes)
        self._fill(slide, 'brand.accessibleBlue')
        self._logo_at(slide, 'white', MARGIN_IN, 0.6, 0.75)
        self._text(slide, title, MARGIN_IN, 2.6, 11.5, 1.8, 44, 'ink.onDark', bold=True, heading=True, anchor='bottom')
        if subtitle:
            self._text(slide, subtitle, MARGIN_IN, 4.55, 11.5, 0.9, 20, 'ink.accentOnBrand')
        return slide

    def section_slide(self, title, kicker=None, notes=None):
        slide = self._new(notes)
        self._fill(slide, 'brand.accessibleBlue')
        if kicker:
            self._text(slide, kicker, MARGIN_IN, 2.35, 11.5, 0.6, 18, 'ink.accentOnBrand')
        self._text(slide, title, MARGIN_IN, 2.95, 11.5, 1.6, 40, 'ink.onDark', bold=True, heading=True)
        self._logo_at(slide, 'white', SLIDE_W_IN - MARGIN_IN - 1.0, SLIDE_H_IN - 0.6, 0.42)
        return slide

    def bullets_slide(self, heading, bullets, notes=None, size=20):
        """bullets: strings, or (text, level) tuples for one level of nesting."""
        slide = self._new(notes)
        self._chrome(slide, heading)
        self._bullet_box(slide, bullets, MARGIN_IN, 1.6, SLIDE_W_IN - 2 * MARGIN_IN, 5.0, size)
        return slide

    def _bullet_box(self, slide, bullets, x, y, w, h, size):
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        for index, item in enumerate(bullets):
            text, level = (item, 0) if isinstance(item, str) else item
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.level = level
            paragraph.space_after = Pt(10)
            _bullet(paragraph, '–' if level else '•', self.kb.hex('brand.cyanBlue'), indent_in=0.3 + 0.35 * level)
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(size - 2 * level)
            run.font.name = self.fonts['body']
            run.font.color.rgb = self._rgb('ink.primary')
        return box

    def two_column_slide(self, heading, left, right, notes=None):
        """left/right: a list of bullet strings, or a path to an image."""
        slide = self._new(notes)
        self._chrome(slide, heading)
        col_w = (SLIDE_W_IN - 2 * MARGIN_IN - 0.5) / 2
        for index, content in enumerate((left, right)):
            x = MARGIN_IN + index * (col_w + 0.5)
            if isinstance(content, str) and os.path.exists(content):
                self._picture(slide, content, x, 1.6, col_w, 5.0)
            else:
                items = content if isinstance(content, list) else [content]
                self._bullet_box(slide, items, x, 1.6, col_w, 5.0, 18)
        return slide

    def kpi_slide(self, heading, kpis, notes=None):
        """kpis: up to four (value, label) pairs, e.g. ('84%', 'OEE, Q3')."""
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        slide = self._new(notes)
        self._chrome(slide, heading)
        count = max(1, min(len(kpis), 4))
        gap = 0.35
        card_w = (SLIDE_W_IN - 2 * MARGIN_IN - gap * (count - 1)) / count
        for index, (value, label) in enumerate(kpis[:4]):
            x = MARGIN_IN + index * (card_w + gap)
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.0), Inches(card_w), Inches(2.8))
            card.fill.solid()
            card.fill.fore_color.rgb = self._rgb('surface.page')
            card.line.color.rgb = self._rgb('surface.border')
            # Drop the theme style reference: it adds a shadow and theme line/fill LibreOffice
            # still honours even with an explicit fill. The explicit fill and line stay.
            style = card._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
            if style is not None:
                card._element.remove(style)
            self._text(slide, value, x + 0.3, 2.35, card_w - 0.6, 1.3, 48, 'ink.heading', bold=True, heading=True)
            self._text(slide, label, x + 0.3, 3.75, card_w - 0.6, 0.8, 16, 'ink.muted')
        return slide

    def chart_slide(self, heading, categories, series, kind='column', notes=None, number_format='General',
                    legend=None):
        """A native, editable PowerPoint chart. series: {'name': [values…]} in display order."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
        from pptx.util import Inches, Pt

        slide = self._new(notes)
        self._chrome(slide, heading)
        data = CategoryChartData()
        data.categories = categories
        for name, values in series.items():
            data.add_series(name, values)
        chart_type = {
            'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
            'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE,
            'stacked': XL_CHART_TYPE.COLUMN_STACKED,
        }[kind]
        frame = slide.shapes.add_chart(chart_type, Inches(MARGIN_IN), Inches(1.55), Inches(SLIDE_W_IN - 2 * MARGIN_IN), Inches(5.2), data)
        chart = frame.chart
        chart.has_title = False  # the slide heading is the title; a one-series chart adds its own
        chart.font.size = Pt(12)
        chart.font.name = self.fonts['body']
        chart.font.color.rgb = self._rgb('ink.secondary')
        show_legend = legend if legend is not None else (len(series) > 1 or kind == 'pie')
        chart.has_legend = show_legend
        if show_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        palette = self.kb.KNAUF['categorical']
        plot = chart.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size = Pt(11)
        labels.font.color.rgb = self._rgb('ink.primary')
        labels.number_format, labels.number_format_is_linked = number_format, False
        if kind == 'pie':
            for index, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self._rgb(f'categorical.{index % len(palette)}')
        else:
            if kind == 'line':
                labels.position = XL_LABEL_POSITION.ABOVE
            else:
                # outside-end labels corrupt a stacked chart; inside-end is the safe choice there
                labels.position = XL_LABEL_POSITION.INSIDE_END if kind == 'stacked' else XL_LABEL_POSITION.OUTSIDE_END
                plot.gap_width = 80
            for index, chart_series in enumerate(plot.series):
                token = f'categorical.{index % len(palette)}'
                if kind == 'line':
                    chart_series.format.line.color.rgb = self._rgb(token)
                    chart_series.smooth = False
                else:
                    chart_series.format.fill.solid()
                    chart_series.format.fill.fore_color.rgb = self._rgb(token)
            value_axis = chart.value_axis
            if kind != 'line':
                value_axis.minimum_scale = 0  # bar lengths only compare honestly from zero
            value_axis.has_major_gridlines = True
            value_axis.major_gridlines.format.line.color.rgb = self._rgb('surface.neutral')
            value_axis.format.line.fill.background()
            value_axis.tick_labels.font.color.rgb = self._rgb('ink.muted')
            chart.category_axis.tick_labels.font.color.rgb = self._rgb('ink.muted')
            chart.category_axis.format.line.color.rgb = self._rgb('surface.neutral')
        return slide

    def table_slide(self, heading, rows, notes=None, col_widths_in=None):
        from pptx.util import Inches, Pt

        slide = self._new(notes)
        self._chrome(slide, heading)
        ncols = max(len(r) for r in rows)
        width = SLIDE_W_IN - 2 * MARGIN_IN
        shape = slide.shapes.add_table(len(rows), ncols, Inches(MARGIN_IN), Inches(1.6), Inches(width), Inches(0.45 * len(rows)))
        table = shape.table
        widths = col_widths_in or [width / ncols] * ncols
        for c, w in enumerate(widths):
            table.columns[c].width = Inches(w)
        for r, row in enumerate(rows):
            for c in range(ncols):
                cell = table.cell(r, c)
                cell.text = str(row[c]) if c < len(row) else ''
                cell.fill.solid()
                cell.fill.fore_color.rgb = self._rgb('surface.neutral' if r == 0 else ('surface.page' if r % 2 == 0 else 'surface.card'))
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(14)
                        run.font.bold = r == 0
                        run.font.name = self.fonts['body']
                        run.font.color.rgb = self._rgb('ink.primary')
        return slide

    def image_slide(self, heading, image_path, caption=None, notes=None):
        slide = self._new(notes)
        self._chrome(slide, heading)
        bottom = 6.1 if caption else 6.7
        self._picture(slide, image_path, MARGIN_IN, 1.55, SLIDE_W_IN - 2 * MARGIN_IN, bottom - 1.55)
        if caption:
            self._text(slide, caption, MARGIN_IN, 6.2, SLIDE_W_IN - 2 * MARGIN_IN, 0.4, 12, 'ink.muted')
        return slide

    def statement_slide(self, statement, attribution=None, notes=None):
        slide = self._new(notes)
        self._fill(slide, 'surface.page')
        self._text(slide, statement, 1.2, 1.8, SLIDE_W_IN - 2.4, 3.2, 34, 'ink.heading', bold=True, heading=True, anchor='middle')
        if attribution:
            self._text(slide, attribution, 1.2, 5.2, SLIDE_W_IN - 2.4, 0.5, 16, 'ink.muted')
        self._logo_at(slide, 'logo', SLIDE_W_IN - MARGIN_IN - 1.0, SLIDE_H_IN - 0.55, 0.42)
        return slide

    def closing_slide(self, message='Thank you', contact=None, notes=None):
        slide = self.section_slide(message, notes=notes)
        if contact:
            self._text(slide, contact, MARGIN_IN, 4.7, 11.5, 0.8, 18, 'ink.accentOnBrand')
        return slide

    def _picture(self, slide, path, x, y, w, h):
        """Fit an image inside the box without stretching it."""
        from PIL import Image
        from pptx.util import Inches

        with Image.open(path) as image:
            ratio = image.width / image.height
        width, height = (w, w / ratio) if w / ratio <= h else (h * ratio, h)
        slide.shapes.add_picture(path, Inches(x + (w - width) / 2), Inches(y + (h - height) / 2), Inches(width), Inches(height))

    def save(self, path):
        self.prs.save(path)
        return path


def _slide_number_field(paragraph):
    """Turn the paragraph's first run into PowerPoint's slide-number field (a:fld slidenum)."""
    import uuid

    from pptx.oxml.ns import qn

    run = paragraph.runs[0]._r
    field = run.makeelement(qn('a:fld'), {'id': '{' + str(uuid.uuid4()).upper() + '}', 'type': 'slidenum'})
    for child in list(run):
        field.append(child)  # a:rPr then a:t — the order a:fld requires
    run.addprevious(field)
    run.getparent().remove(run)


def _bullet(paragraph, char, colour_hex, indent_in=0.3):
    """A real bullet (a:buChar), never a typed '•' — screen readers and PowerPoint agree."""
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    ppr = paragraph._p.get_or_add_pPr()
    ppr.set('marL', str(int(Inches(indent_in))))
    ppr.set('indent', str(-int(Inches(0.3))))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buClr'):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    colour = ppr.makeelement(qn('a:buClr'), {})
    srgb = colour.makeelement(qn('a:srgbClr'), {'val': colour_hex})
    colour.append(srgb)
    bullet = ppr.makeelement(qn('a:buChar'), {'char': char})
    # buClr must precede buChar (CT_TextParagraphProperties sequence)
    ppr.append(colour)
    ppr.append(bullet)


# ── lint ─────────────────────────────────────────────────────────────────────────────────


def _font_file(family, bold=False):
    style = 'bold' if bold else 'regular'
    try:
        return subprocess.run(['fc-match', '-f', '%{file}', f'{family}:{style}'], capture_output=True,
                              text=True, timeout=10).stdout.strip() or None
    except Exception:
        return None


def _wrapped_height_in(frame, width_in, default_family):
    """Rough height the text needs, from real font metrics (Pillow + fontconfig)."""
    from PIL import ImageFont

    total_pt = 0.0
    for paragraph in frame.paragraphs:
        runs = paragraph.runs
        text = ''.join(r.text for r in runs)
        size = next((r.font.size.pt for r in runs if r.font.size), 18.0)
        family = next((r.font.name for r in runs if r.font.name), default_family)
        bold = any(r.font.bold for r in runs)
        path = _font_file(family, bold)
        try:
            font = ImageFont.truetype(path, int(size * 10)) if path else None
        except Exception:
            font = None
        width_pt = width_in * 72 * 10
        lines, current = 1, ''
        for word in text.split():
            candidate = f'{current} {word}'.strip()
            measured = font.getlength(candidate) if font else len(candidate) * size * 5.2
            if measured > width_pt and current:
                lines, current = lines + 1, word
            else:
                current = candidate
        total_pt += lines * size * 1.2 + (paragraph.space_after.pt if paragraph.space_after else 0)
    return total_pt / 72


def lint(path):
    """Layout problems PowerPoint users would see: off-slide shapes, likely overflow, leftovers."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(path)
    width, height = prs.slide_width, prs.slide_height
    problems = []
    for number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            name = f'slide {number} "{shape.name}"'
            if shape.left is not None and shape.width is not None:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width + Emu(9525) or shape.top + shape.height > height + Emu(9525):
                    problems.append(f'{name} extends beyond the slide edge')
            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                problems.append(f'{name} is an empty placeholder (delete it or fill it)')
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            frame = shape.text_frame
            margins = (frame.margin_left or 0) + (frame.margin_right or 0)
            box_w = (shape.width - margins) / 914400
            box_h = (shape.height - (frame.margin_top or 0) - (frame.margin_bottom or 0)) / 914400
            if box_w <= 0:
                continue
            needed = _wrapped_height_in(frame, box_w, 'Arial')
            if needed > box_h * 1.05 + 0.05:
                problems.append(
                    f'{name}: text needs about {needed:.1f} in but the box is {box_h:.1f} in tall '
                    '(likely overflow — shorten, split the slide, or enlarge the box)'
                )
    return problems


# ── slide operations ─────────────────────────────────────────────────────────────────────


def _slide_ids(prs):
    return prs.slides._sldIdLst  # noqa: SLF001 - python-pptx exposes no public reorder API


def delete_slides(prs, numbers):
    ids = _slide_ids(prs)
    entries = list(ids)
    for number in sorted(set(numbers), reverse=True):
        entry = entries[number - 1]
        prs.part.drop_rel(entry.get(f'{{{R_NS}}}id'))
        ids.remove(entry)


def reorder_slides(prs, order):
    ids = _slide_ids(prs)
    entries = list(ids)
    if sorted(order) != list(range(1, len(entries) + 1)):
        raise SystemExit(f'--order must list every slide 1-{len(entries)} exactly once')
    for entry in entries:
        ids.remove(entry)
    for number in order:
        ids.append(entries[number - 1])


def duplicate_slide(prs, number):
    """Append a copy of slide `number`. Pictures and charts stay shared with the original."""
    source = prs.slides[number - 1]
    copy_slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(copy_slide.shapes):
        shape._element.getparent().remove(shape._element)
    rel_map = {}
    for rel_id, rel in source.part.rels.items():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
            continue
        rel_map[rel_id] = (
            copy_slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            if rel.is_external else copy_slide.part.relate_to(rel.target_part, rel.reltype)
        )
    for element in source.shapes._spTree.iterchildren():
        if element.tag.endswith(('}nvGrpSpPr', '}grpSpPr')):
            continue
        clone = copy.deepcopy(element)
        for node in clone.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith(f'{{{R_NS}}}') and value in rel_map:
                    node.set(attr, rel_map[value])
        copy_slide.shapes._spTree.append(clone)
    if source.has_notes_slide:
        copy_slide.notes_slide.notes_text_frame.text = source.notes_slide.notes_text_frame.text
    background = source._element.cSld.bg
    if background is not None:
        copy_slide._element.cSld.insert(0, copy.deepcopy(background))
    return len(prs.slides)


# ── CLI ──────────────────────────────────────────────────────────────────────────────────


def _save(prs, source, out):
    target = out or source
    if not out and not os.path.abspath(source).startswith('/mnt/data/'):
        raise SystemExit('Give -o: the input is outside /mnt/data and may be read-only.')
    prs.save(target)
    print(f'wrote {target}')


def main(argv=None):
    parser = argparse.ArgumentParser(prog='karl_pptx.py')
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('qa'); p.add_argument('file'); p.add_argument('--pages')
    p = sub.add_parser('lint'); p.add_argument('file')
    p = sub.add_parser('text'); p.add_argument('file')
    p = sub.add_parser('to-pdf'); p.add_argument('file'); p.add_argument('--out')
    p = sub.add_parser('duplicate'); p.add_argument('file'); p.add_argument('--slide', type=int, required=True); p.add_argument('-o', '--out')
    p = sub.add_parser('delete'); p.add_argument('file'); p.add_argument('--slides', required=True); p.add_argument('-o', '--out')
    p = sub.add_parser('reorder'); p.add_argument('file'); p.add_argument('--order', required=True); p.add_argument('-o', '--out')
    args = parser.parse_args(argv)

    if args.command == 'qa':
        office_qa.ensure_fonts()
        return office_qa.qa(args.file, extra_problems=lint(args.file), pages=office_qa.parse_pages(args.pages), label='Slide')
    if args.command == 'lint':
        problems = lint(args.file)
        print(json.dumps({'file': args.file, 'problems': problems}, indent=1))
        return 1 if problems else 0
    if args.command == 'text':
        from pptx import Presentation

        for number, slide in enumerate(Presentation(args.file).slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text)
                elif getattr(shape, 'has_table', False) and shape.has_table:
                    texts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
            print(f'--- slide {number} ---')
            print('\n'.join(texts))
            if notes.strip():
                print(f'[notes] {notes}')
        return 0
    if args.command == 'to-pdf':
        office_qa.ensure_fonts()
        pdf = office_qa.soffice(args.file, 'pdf')
        out = args.out or os.path.splitext(args.file)[0] + '.pdf'
        with open(pdf, 'rb') as src, open(out, 'wb') as dst:
            dst.write(src.read())
        print(f'wrote {out}')
        return 0
    from pptx import Presentation

    prs = Presentation(args.file)
    if args.command == 'duplicate':
        print(f'slide {args.slide} copied to position {duplicate_slide(prs, args.slide)}')
    elif args.command == 'delete':
        delete_slides(prs, [int(n) for n in args.slides.split(',')])
    elif args.command == 'reorder':
        reorder_slides(prs, [int(n) for n in args.order.split(',')])
    _save(prs, args.file, args.out)
    return 0


if __name__ == '__main__':
    sys.exit(main())
